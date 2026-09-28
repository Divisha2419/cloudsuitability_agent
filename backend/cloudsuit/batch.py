"""Batch mode (behaviour rule 7): assess many applications from a CSV/Excel upload."""

from __future__ import annotations

import csv
import io
import re

from openpyxl import load_workbook

from .engine import assess
from .schema import fields, validate


class BatchError(ValueError):
    pass


def _header_key(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\(\*\)", "", str(text or ""))).strip().casefold()


def _column_map(headers: list[str]) -> dict[int, str]:
    """Map column index -> field id. Headers may be field ids or field labels."""
    lookup = {}
    for fid, f in fields().items():
        lookup[fid.casefold()] = fid
        lookup[_header_key(f["label"])] = fid
    mapping = {i: lookup[_header_key(h)] for i, h in enumerate(headers) if _header_key(h) in lookup}
    if "app_name" not in mapping.values():
        raise BatchError('No "Application Name" (or "app_name") column found in the header row.')
    return mapping


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = [["" if v is None else str(v) for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise BatchError("Upload a .csv or .xlsx file.")
    if not rows:
        raise BatchError("The file is empty.")
    mapping = _column_map(rows[0])
    out = []
    for r in rows[1:]:
        record = {fid: r[i].strip() for i, fid in mapping.items() if i < len(r) and r[i] and r[i].strip()}
        if record:
            out.append(record)
    return out


def assess_batch(filename: str, content: bytes) -> list[dict]:
    """Returns one entry per data row: {row, answers, errors, result}."""
    entries = []
    for n, record in enumerate(read_rows(filename, content), start=2):
        answers, errors = validate(record)
        entries.append({"row": n, "answers": answers, "errors": errors, "result": assess(record)})
    return entries


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f["label"] for f in fields().values()])
    return buf.getvalue()
