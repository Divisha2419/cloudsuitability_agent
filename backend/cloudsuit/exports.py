"""PDF and Excel exports of single assessments and of the portfolio."""

from __future__ import annotations

import io
from xml.sax.saxutils import escape as e

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .schema import fields, option_label

RATING_FILL = {"cloud_ready": "D9F2E3", "needs_upgrade": "FFF1CC", "not_suitable": "F9D6D5", "na": "EEEEEE"}
SEVERITY_FILL = {"high": "F9D6D5", "medium": "FFF1CC", "low": "E3ECF7"}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


# --------------------------------------------------------------------------- Excel


def _header(ws, row: list[str]) -> None:
    ws.append(row)
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT


def _autosize(ws, max_width: int = 80) -> None:
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(12, width + 2), max_width)
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def assessment_xlsx(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    app, rec, p3 = result["application"], result["recommendation"], result["phase3"]
    rows = [
        ("Application", app["name"]), ("Application ID", app["id"]), ("Assessed by", app["manager"]),
        ("Date", app["date"]), ("Phase 1 – Hard Filter", result["phase1"]["status"]),
        *[("Hard filter reason", f["reason"]) for f in result["phase1"]["filters"]],
        ("Phase 2 – Tech Stack", result["phase2"]["overall_label"]),
        ("Cloud Native Score", f"{p3['total']}/{p3['max']}"), ("Readiness Band", p3["band"]["label"]),
        ("6R Recommendation", rec["headline"]), ("Definition", rec["definition"]),
        *[("Rationale", r) for r in rec["rationale"]], *[("Note", n) for n in rec["notes"]],
        *[("Next step", s) for s in rec["next_steps"]],
    ]
    _header(ws, ["Item", "Value"])
    for r in rows:
        ws.append(list(r))
    _autosize(ws)

    ws = wb.create_sheet("Phase 2 Tech Stack")
    _header(ws, ["Component", "Input", "Rating", "Matched rule", "Detail"])
    for c in result["phase2"]["components"]:
        ws.append([c["label"], c["input"], c["rating_label"], c["matched"], c["detail"]])
        ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=RATING_FILL[c["rating"]])
    ws.append(["Overall", "", result["phase2"]["overall_label"], "", ""])
    _autosize(ws)

    ws = wb.create_sheet("Phase 3 Score")
    _header(ws, ["Dimension", "Answer", "Score", "Weight"])
    for d in p3["dimensions"]:
        ws.append([d["label"], d["answer"], d["score"], d["weight"]])
    ws.append(["Total", p3["band"]["label"], p3["total"], p3["max"]])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    _autosize(ws)

    ws = wb.create_sheet("Risks & Flags")
    _header(ws, ["Severity", "Risk / Flag"])
    for r in result["risks"]:
        ws.append([r["severity"].title(), r["message"]])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=SEVERITY_FILL[r["severity"]])
    _autosize(ws)

    ws = wb.create_sheet("Answers")
    _header(ws, ["Field", "Answer"])
    for fid, field in fields().items():
        if fid in result["answers"]:
            ws.append([field["label"], option_label(fid, result["answers"][fid])])
    _autosize(ws)
    return _to_bytes(wb)


PORTFOLIO_COLUMNS = ["Rank", "Application", "ID", "Manager", "Score", "Band", "Phase 1", "Tech Stack", "6R Recommendation"]


def portfolio_rows(results: list[dict]) -> list[list]:
    ranked = sorted(results, key=lambda r: r["phase3"]["total"], reverse=True)
    return [
        [
            i, r["application"]["name"], r["application"]["id"], r["application"]["manager"],
            r["phase3"]["total"], r["phase3"]["band"]["label"], r["phase1"]["status"],
            r["phase2"]["overall_label"], r["recommendation"]["headline"],
        ]
        for i, r in enumerate(ranked, 1)
    ]


def portfolio_xlsx(results: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Portfolio"
    _header(ws, PORTFOLIO_COLUMNS)
    for row in portfolio_rows(results):
        ws.append(row)
    _autosize(ws)

    ws = wb.create_sheet("6R Summary")
    _header(ws, ["6R Recommendation", "Applications"])
    counts: dict[str, int] = {}
    for r in results:
        counts[r["recommendation"]["headline"]] = counts.get(r["recommendation"]["headline"], 0) + 1
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        ws.append([k, v])
    _autosize(ws)
    return _to_bytes(wb)


def batch_template_xlsx() -> bytes:
    """Upload template: one column per attribute, dropdowns for select fields."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    all_fields = list(fields().values())
    _header(ws, [f["label"] + (" (*)" if f.get("required") else "") for f in all_fields])
    for i, f in enumerate(all_fields, 1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = max(16, len(f["label"]) + 4)
        if f["type"] == "select":
            values = ",".join(str(o["value"]) for o in f["options"])
            dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}500")
    ref = wb.create_sheet("Allowed values")
    _header(ref, ["Column", "Field ID", "Required", "Allowed values / example", "Shown only when"])
    for f in all_fields:
        allowed = " | ".join(str(o["value"]) for o in f.get("options", [])) or f.get("placeholder", "free text")
        cond = f.get("show_if")
        ref.append([
            f["label"], f["id"], "Yes" if f.get("required") else "No", allowed,
            f"{fields()[cond['field']]['label']} is {' or '.join(cond['in'])}" if cond else "",
        ])
    _autosize(ref)
    return _to_bytes(wb)


# --------------------------------------------------------------------------- PDF


def _pdf_color(hex_: str):
    return colors.HexColor("#" + hex_)


def assessment_pdf(result: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"Cloud Suitability – {result['application']['name']}")
    ss = getSampleStyleSheet()
    body = ss["BodyText"]
    h2 = ss["Heading2"]
    app, p1, p2, p3, rec = (result[k] for k in ("application", "phase1", "phase2", "phase3", "recommendation"))
    story = [
        Paragraph("Cloud Suitability Assessment Report", ss["Title"]),
        Paragraph(f"<b>Application:</b> {e(app['name'] or '—')} &nbsp;|&nbsp; <b>ID:</b> {e(app['id'] or '—')}", body),
        Paragraph(f"<b>Assessed by:</b> {e(app['manager'] or '—')} &nbsp;|&nbsp; <b>Date:</b> {app['date']}", body),
        Spacer(1, 6),
    ]

    def table(data, widths, extra=()):
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _pdf_color("1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            *extra,
        ]))
        return t

    def p(text):
        return Paragraph(text, body)

    # Recommendation first: it is what readers look for.
    story += [
        Paragraph("6R Recommendation", h2),
        table([["Recommendation", "Cloud Native Score", "Band"],
               [p(f"<b>★ {e(rec['headline'])}</b>"), f"{p3['total']}/{p3['max']}", p(e(p3["band"]["label"]))]],
              [70 * mm, 40 * mm, 64 * mm]),
        Spacer(1, 4), p(e(rec["definition"])), Spacer(1, 4), p("<b>Rationale</b>"),
        *[p(f"• {e(r)}") for r in rec["rationale"]], *[p(f"<b>!</b> {e(n)}") for n in rec["notes"]],
        Spacer(1, 4), p("<b>Next steps</b>"), *[p(f"{i}. {e(s)}") for i, s in enumerate(rec["next_steps"], 1)],
    ]

    story += [Paragraph("Phase 1 – Hard Filter Result", h2), p(f"<b>{p1['status']}</b>")]
    story += [p(f"• {e(f['reason'])} → {e(f['outcome'])}") for f in p1["filters"]]
    if p1["triggered"]:
        story.append(p("<i>Phases 2 and 3 are shown for roadmap planning only.</i>"))

    rows = [["Component", "Input", "Rating", "Matched rule"]]
    styles = []
    for i, c in enumerate(p2["components"], 1):
        rows.append([c["label"], p(e(c["input"])), c["rating_label"], p(e(c["matched"]))])
        styles.append(("BACKGROUND", (2, i), (2, i), _pdf_color(RATING_FILL[c["rating"]])))
    rows.append(["Overall", "", p2["overall_label"], ""])
    styles.append(("FONTNAME", (0, len(rows) - 1), (-1, len(rows) - 1), "Helvetica-Bold"))
    story += [Paragraph("Phase 2 – Tech Stack Cloud Suitability", h2),
              table(rows, [36 * mm, 52 * mm, 36 * mm, 50 * mm], styles)]

    rows = [["Dimension", "Answer", "Score"]]
    rows += [[d["label"], p(e(d["answer"])), f"{d['score']}/{d['weight']}"] for d in p3["dimensions"]]
    rows.append(["Total", p3["band"]["label"], f"{p3['total']}/{p3['max']}"])
    story += [Paragraph("Phase 3 – Cloud Native Score", h2),
              table(rows, [70 * mm, 76 * mm, 28 * mm],
                    [("FONTNAME", (0, len(rows) - 1), (-1, len(rows) - 1), "Helvetica-Bold")])]

    story.append(Paragraph("Key Risks & Flags", h2))
    if result["risks"]:
        rows = [["Severity", "Risk / Flag"]] + [[r["severity"].title(), p(e(r["message"]))] for r in result["risks"]]
        styles = [("BACKGROUND", (0, i), (0, i), _pdf_color(SEVERITY_FILL[r["severity"]]))
                  for i, r in enumerate(result["risks"], 1)]
        story.append(table(rows, [24 * mm, 150 * mm], styles))
    else:
        story.append(p("None identified."))

    doc.build(story)
    return buf.getvalue()
