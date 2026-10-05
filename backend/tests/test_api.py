import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from cloudsuit import api
from cloudsuit.storage import Store


@pytest.fixture
def client(tmp_path, monkeypatch):
    s = Store(f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setattr(api, "store", lambda: s)
    return TestClient(api.app)


@pytest.fixture
def admin(client):
    token = client.post("/api/admin/login", json={"username": "admin", "password": "admin"}).json()["token"]
    return {"Authorization": f"Bearer {token}"}


def test_schema_lists_steps_sections_and_projects(client):
    schema = client.get("/api/schema").json()
    assert [s["id"] for s in schema["steps"]] == ["application_info", "business", "technical", "results"]
    project = schema["sections"][0]["fields"][0]
    assert project["id"] == "project" and project["required"]
    assert [o["value"] for o in project["options"]] == ["ABB Edge China", "SDD", "Indian Bank"]
    assert client.get("/api/projects").json() == ["ABB Edge China", "SDD", "Indian Bank"]


def test_assess_accepts_partial_answers(client):
    r = client.post("/api/assess", json={"answers": {"architecture": "SOA"}}).json()
    assert r["complete"] is False
    assert r["phase3"]["total"] >= 10


def test_submit_requires_complete_answers(client):
    r = client.post("/api/assessments", json={"answers": {"app_name": "X"}})
    assert r.status_code == 422
    assert {"architecture", "project"} <= set(r.json()["detail"]["errors"])


def test_unknown_project_is_rejected(client, answers):
    r = client.post("/api/assessments", json={"answers": {**answers, "project": "Not a project"}})
    assert r.status_code == 422 and "project" in r.json()["detail"]["errors"]


def test_admin_endpoints_need_login(client):
    assert client.post("/api/admin/login", json={"username": "admin", "password": "wrong"}).status_code == 401
    for path in ("/api/admin/projects", "/api/admin/assessments?project=SDD", "/api/admin/assessments/1",
                 "/api/admin/export.xlsx?project=SDD"):
        assert client.get(path).status_code == 401, path
    assert client.delete("/api/admin/assessments/1").status_code == 401
    assert client.get("/api/admin/projects", headers={"Authorization": "Bearer nonsense"}).status_code == 401


def test_submit_replaces_same_app_id_in_same_project(client, admin, answers):
    first = client.post("/api/assessments", json={"answers": answers}).json()
    assert first["result"]["cloud_suitability"] == {"suitable": True, "label": "Cloud Suitable"}
    second = client.post("/api/assessments", json={"answers": {**answers, "operating_system": "Windows Server 2012"}}).json()
    assert second["id"] == first["id"]
    # Same Application ID in another project is a separate application.
    other = client.post("/api/assessments", json={"answers": {**answers, "project": "SDD"}}).json()
    assert other["id"] != first["id"]

    abb = client.get("/api/admin/assessments", params={"project": "ABB Edge China"}, headers=admin).json()
    assert abb["total"] == 1 and abb["rows"][0]["recommendation"] == "replatform"
    projects = client.get("/api/admin/projects", headers=admin).json()
    assert projects == [{"name": "ABB Edge China", "count": 1}, {"name": "SDD", "count": 1}, {"name": "Indian Bank", "count": 0}]


def test_admin_table_summary_detail_delete(client, admin, answers):
    client.post("/api/assessments", json={"answers": answers})
    retain = {**answers, "app_id": "APP-002", "app_name": "Line SCADA", "latency": "Ultra Low"}
    client.post("/api/assessments", json={"answers": retain})

    data = client.get("/api/admin/assessments", params={"project": "ABB Edge China"}, headers=admin).json()
    assert (data["total"], data["suitable"], data["not_suitable"]) == (2, 1, 1)
    assert data["six_r"] == {"rehost": 1, "retain": 1}
    ok, no = data["rows"]
    assert [ok["s_no"], ok["app_id"], ok["suitability"], ok["recommendation_headline"]] == [1, "APP-001", "Cloud Suitable", "REHOST"]
    assert ok["score"] == 95 and 1 <= len(ok["rationale"]) <= 3
    assert no["suitability"] == "Not Cloud Suitable" and no["score"] is None

    detail = client.get(f"/api/admin/assessments/{no['id']}", headers=admin).json()
    assert detail["result"]["recommendation"]["code"] == "retain"
    assert detail["answers"]["app_name"] == "Line SCADA"

    assert client.delete(f"/api/admin/assessments/{no['id']}", headers=admin).status_code == 204
    assert client.get(f"/api/admin/assessments/{no['id']}", headers=admin).status_code == 404
    assert client.get("/api/admin/assessments", params={"project": "ABB Edge China"}, headers=admin).json()["total"] == 1


def test_logout_invalidates_token(client, admin):
    assert client.post("/api/admin/logout", headers=admin).status_code == 204
    assert client.get("/api/admin/projects", headers=admin).status_code == 401


def test_exports(client, admin, answers):
    pdf = client.post("/api/report/pdf", json={"answers": {**answers, "app_name": "A & B <test>"}})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    xlsx = client.post("/api/report/xlsx", json={"answers": answers})
    assert load_workbook(io.BytesIO(xlsx.content)).sheetnames[0] == "Summary"
    client.post("/api/assessments", json={"answers": answers})
    wb = load_workbook(io.BytesIO(client.get("/api/admin/export.xlsx", params={"project": "ABB Edge China"}, headers=admin).content))
    ws = wb["Applications"]
    assert [c.value for c in ws[1]][:4] == ["S.No", "Application ID", "Application Name", "Cloud Suitability Result"]
    assert ws["C2"].value == "Order Portal" and ws["F2"].value == 95


def test_old_database_without_project_column_is_upgraded(tmp_path, answers):
    url = f"sqlite:///{tmp_path / 'old.db'}"
    eng = create_engine(url)
    with eng.begin() as c:
        c.execute(text("CREATE TABLE assessments (id INTEGER PRIMARY KEY, app_name VARCHAR(255), app_id VARCHAR(255),"
                       " answers JSON, result JSON, created_at DATETIME, updated_at DATETIME)"))
    from cloudsuit.engine import assess

    old = assess({k: v for k, v in answers.items() if k != "project"})
    with eng.begin() as c:
        c.execute(text("INSERT INTO assessments (app_name, app_id, answers, result, created_at, updated_at) "
                       "VALUES ('Old', 'OLD-1', '{}', :r, '2026-01-01', '2026-01-01')"), {"r": __import__("json").dumps(old)})
    s = Store(url)
    assert s.project_counts() == {"Unassigned": 1}
    assert s.list("Unassigned")[0].row()["suitability"] == "Cloud Suitable"


def test_logo_404_without_file_and_served_with_file(client, tmp_path, monkeypatch):
    monkeypatch.setattr(api, "BRANDING_DIR", tmp_path)  # independent of the real config/branding
    assert client.get("/api/branding/logo").status_code == 404
    assert client.get("/api/branding").json() == {"logo_url": None}
    (tmp_path / "logo.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>")
    r = client.get("/api/branding/logo")
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/svg+xml")
    assert client.get("/api/branding").json() == {"logo_url": "/api/branding/logo"}


def test_wave_roadmap_only_above_five_applications(client, admin, answers):
    def submit(app_id, **over):
        client.post("/api/assessments", json={"answers": {**answers, "app_id": app_id, "app_name": app_id, **over}})

    for i in range(5):
        submit(f"A{i}")
    data = client.get("/api/admin/assessments", params={"project": "ABB Edge China"}, headers=admin).json()
    assert data["total"] == 5 and data["roadmap"] is None

    # A0–A4: rehost, High criticality (business critical) -> Wave 2
    submit("LOW", business_criticality="Low")                                   # rehost, low -> Wave 1
    submit("RPL", business_criticality="Low", operating_system="Windows Server 2012")  # replatform -> Wave 2
    submit("RET", latency="Ultra Low")                                          # retain -> out of scope
    roadmap = client.get("/api/admin/assessments", params={"project": "ABB Edge China"}, headers=admin).json()["roadmap"]
    by_wave = {w["name"]: [a["app_id"] for a in w["applications"]] for w in roadmap["waves"]}
    assert by_wave["Wave 1 – Quick wins"] == ["LOW"]
    assert by_wave["Wave 2 – Core migration"][0] == "RPL"  # Low criticality before High
    assert set(by_wave["Wave 2 – Core migration"]) == {"RPL", "A0", "A1", "A2", "A3", "A4"}
    assert [a["app_id"] for a in roadmap["out_of_scope"]["applications"]] == ["RET"]
    assert "Wave 3 – SaaS replacement" not in by_wave  # empty waves are left out

    wb = load_workbook(io.BytesIO(client.get("/api/admin/export.xlsx", params={"project": "ABB Edge China"}, headers=admin).content))
    assert "Wave Roadmap (provisional)" in wb.sheetnames
